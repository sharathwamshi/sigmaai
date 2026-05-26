import os
import json
import openpyxl

BASE_DIR = os.path.dirname(__file__)

XLSX_PATH = os.path.join(BASE_DIR, "sigma_products.xlsx")
PRICES_PATH = os.path.join(BASE_DIR, "data", "prices.json")

SHEETS = [
    "Still Lenses",
    "Cine Lenses",
    "Cameras",
    "Accessories"
]

prices = {}

wb = openpyxl.load_workbook(XLSX_PATH, read_only=True, data_only=True)

for sheet_name in SHEETS:

    if sheet_name not in wb.sheetnames:
        continue

    ws = wb[sheet_name]

    rows = list(ws.iter_rows(values_only=True))

    if not rows:
        continue

    headers = [str(h).strip() if h else "" for h in rows[0]]

    try:
        product_col = headers.index("Product Name")
    except:
        continue

    for row in rows[1:]:

        if not row:
            continue

        name = row[product_col]

        if not name:
            continue

        name = str(name).strip()

        if not name or name == "None":
            continue

        prices[name] = {
            "price_inr": "₹0",
            "price_range_inr": "₹0 - ₹0",
            "price_usd": 0
        }

os.makedirs(os.path.dirname(PRICES_PATH), exist_ok=True)

with open(PRICES_PATH, "w", encoding="utf-8") as f:
    json.dump(prices, f, indent=2, ensure_ascii=False)

print(f"Generated prices.json with {len(prices)} products")